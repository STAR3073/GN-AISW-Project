from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from uuid import uuid4

import pandas as pd
from docx import Document as DocxDocument

from app.agent.evaluator import evaluate
from app.agent.graph import build_graph
from app.agent.nodes.planner import build_work_plan
from app.agent.nodes.responder import fallback_reply
from app.agent.nodes.tool_executor import tool_executor
from app.agent.nodes.tool_selector import tool_selector
from app.agent.nodes.verifier import verifier
from app.agent.router import after_verify
from app.agent.semantic_intent import resolve_hybrid_intent
from app.tools.excel_analyzer import inspect_and_analyze
from app.tools.registry import invoke_tool, is_registered
from app.tools.report_generator import BANNED_CAUSES, generate_report


def _xlsx(path: Path, rows: list[dict], sheet: str = "품질") -> Path:
    pd.DataFrame(rows).to_excel(path, index=False, sheet_name=sheet)
    return path


def _quality_rows() -> list[dict]:
    return [
        {"품목": "BRACKET", "공정": "용접", "설비": "2호기", "생산일": "2026-08-03", "생산수량": 100, "불량수량": 2, "불량유형": "기공"},
        {"품목": "BRACKET", "공정": "용접", "설비": "2호기", "생산일": "2026-09-03", "생산수량": 100, "불량수량": 12, "불량유형": "기공"},
        {"품목": "HOUSING", "공정": "가공", "설비": "1호기", "생산일": "2026-08-03", "생산수량": 80, "불량수량": 1, "불량유형": "치수"},
        {"품목": "HOUSING", "공정": "가공", "설비": "1호기", "생산일": "2026-09-03", "생산수량": 80, "불량수량": 1, "불량유형": "외관"},
    ]


def _excel_doc(doc_id="doc-1", name="quality.xlsx"):
    return {
        "id": doc_id,
        "filename": name,
        "stored_path": f"/uploads/{name}",
        "raw_text": f"[tabular file: {name}]",
    }


def _report_plan():
    return {
        "goal": "이번 주 품질 현황 보고서 작성",
        "requires_tools": True,
        "steps": [
            {"id": 1, "tool": "analyze_excel", "args": {"document_id": "doc-1"}, "document_id": "doc-1"},
            {"id": 2, "tool": "generate_report", "args": {"depends_on": "analyze_excel"}, "document_id": "doc-1"},
        ],
        "route": "tools",
    }


def _state(message="이번 주 품질 현황 보고서 만들어줘", excel=None, **extra):
    docs = extra.pop("documents", [_excel_doc()])
    results = extra.pop("execution_results", None)
    if results is None:
        results = []
        if excel is not None:
            results.append(
                {
                    "tool": "analyze_excel",
                    "ok": bool(excel.get("ok", True)),
                    "detail": excel.get("detail") or "",
                    "data": excel,
                }
            )
    state = {
        "intent": extra.pop("intent", "execute"),
        "user_message": message,
        "interpreted_goal": extra.pop("interpreted_goal", message),
        "session_id": extra.pop("session_id", ""),
        "documents": docs,
        "document": docs[0] if docs else {},
        "plan": extra.pop("plan", _report_plan()),
        "plan_steps": extra.pop("plan_steps", _report_plan()["steps"]),
        "tool_calls": extra.pop("tool_calls", [{"tool": "analyze_excel", "args": {"document_id": "doc-1"}}]),
        "execution_results": results,
        "retry_count": extra.pop("retry_count", 0),
        "evaluation_count": 0,
        "validated_facts": extra.pop("validated_facts", {}),
        "actions": [],
        "pending_clarification": extra.pop("pending_clarification", None),
    }
    state.update(extra)
    return state


class GenerateReportToolTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.addCleanup(self._tmp.cleanup)
        self.upload_dir = Path(self._tmp.name)

    def _analyze(self, rows):
        path = _xlsx(self.upload_dir / f"quality-{uuid4().hex}.xlsx", rows)
        return inspect_and_analyze(path, filename="quality.xlsx", document_id="doc-1")

    def _generate(self, excel, facts=None, args=None, **extra):
        settings = SimpleNamespace(upload_dir=self.upload_dir)
        with patch("app.tools.report_generator.get_settings", return_value=settings):
            result = generate_report(
                args or {},
                _state(excel=excel, validated_facts=facts or {}, **extra),
            )
        stored = Path(result["stored_path"]) if result.get("stored_path") else None
        text = ""
        if stored and stored.exists():
            doc = DocxDocument(stored)
            paras = [p.text for p in doc.paragraphs]
            cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
            text = "\n".join(paras + cells)
        return result, text

    def test_registered_and_executor_has_no_special_case(self):
        self.assertTrue(is_registered("generate_report"))
        src = Path(__file__).resolve().parents[1] / "app" / "agent" / "nodes" / "tool_executor.py"
        self.assertNotIn("generate_report", src.read_text(encoding="utf-8"))
        self.assertNotIn("elif", src.read_text(encoding="utf-8"))

    def test_e2e_plan_analyze_continue_report_docx(self):
        excel = self._analyze(_quality_rows())
        self.assertTrue(excel["ok"])
        self.assertFalse(excel.get("needs_clarification"))

        intent_state = _state("이번 주 품질 현황 보고서 만들어줘", execution_results=[])
        resolved = resolve_hybrid_intent(intent_state)
        intent_state.update(resolved)
        plan = build_work_plan(intent_state)
        self.assertEqual([s.tool for s in plan.steps], ["analyze_excel", "generate_report"])

        first = tool_selector({**intent_state, "plan": plan.model_dump()})
        self.assertEqual([c["tool"] for c in first["tool_calls"]], ["analyze_excel"])

        excel_state = _state(excel=excel, plan=plan.model_dump(), plan_steps=[s.model_dump() for s in plan.steps])
        first_eval = evaluate(excel_state)
        self.assertEqual(first_eval.status, "continue")
        self.assertIn("total_production", first_eval.validated_facts)
        verified = verifier({**excel_state, "retry_count": 0})
        self.assertEqual(verified["evaluation_status"], "continue")
        self.assertEqual(verified["retry_count"], 0)
        self.assertEqual(after_verify(verified), "continue_tools")

        next_calls = tool_selector({**excel_state, "validated_facts": first_eval.validated_facts})
        self.assertEqual([c["tool"] for c in next_calls["tool_calls"]], ["generate_report"])
        self.assertEqual(next_calls["tool_calls"][0]["args"].get("depends_on"), "analyze_excel")

        report, text = self._generate(excel, facts=first_eval.validated_facts)
        self.assertTrue(report["ok"])
        self.assertTrue(Path(report["stored_path"]).exists())
        self.assertTrue(report["filename"].endswith(".docx"))
        self.assertIn("분석 개요", report["sections_included"])
        self.assertIn("주요 확인사항", report["sections_included"])
        self.assertEqual(report["used_facts"]["total_production"], first_eval.validated_facts["total_production"])
        self.assertEqual(report["used_facts"]["total_defects"], first_eval.validated_facts["total_defects"])
        self.assertIn("주간 품질 현황 보고서", text)
        self.assertIn("생성일시", text)
        self.assertIn("quality.xlsx", text)

        report_state = _state(
            excel=excel,
            validated_facts=first_eval.validated_facts,
            execution_results=[
                {"tool": "analyze_excel", "ok": True, "detail": "", "data": excel},
                {"tool": "generate_report", "ok": True, "detail": report["detail"], "data": report},
            ],
            tool_calls=[{"tool": "generate_report", "args": {"depends_on": "analyze_excel"}}],
        )
        second_eval = evaluate(report_state)
        self.assertEqual(second_eval.status, "pass")
        names = {c.name: c.passed for c in second_eval.checks}
        self.assertTrue(names.get("report_created"))
        self.assertTrue(names.get("report_file_exists"))
        self.assertTrue(names.get("required_sections"))
        self.assertTrue(names.get("report_numbers_match_facts"))
        reply = fallback_reply({**report_state, "evaluation_status": "pass"})
        self.assertIn("품질 현황 보고서가 생성되었습니다", reply)
        self.assertIn(report["filename"], reply)
        self.assertIn("다운로드", reply)
        compiled = build_graph()
        self.assertIn("verifier", compiled.nodes)
        self.assertIn("tool_selector", compiled.nodes)

    def test_missing_production_does_not_invent_numbers(self):
        excel = self._analyze([{"품목": "BRACKET", "불량수량": 4}, {"품목": "HOUSING", "불량수량": 2}])
        eval_result = evaluate(_state(excel=excel))
        self.assertEqual(eval_result.status, "continue")
        self.assertNotIn("total_production", eval_result.validated_facts)
        self.assertNotIn("defect_rate", eval_result.validated_facts)
        report, text = self._generate(excel, facts=eval_result.validated_facts)
        self.assertTrue(report["ok"])
        self.assertNotIn("total_production", report["used_facts"])
        self.assertNotIn("defect_rate", report["used_facts"])
        self.assertIn("생산수량 컬럼이 없어", text)
        self.assertNotRegex(text, r"총 생산수량\s*[:=]\s*\d")
        for banned in BANNED_CAUSES:
            self.assertNotIn(banned, text)

    def test_missing_equipment_omits_equipment_section(self):
        rows = [
            {"품목": "BRACKET", "생산수량": 100, "불량수량": 4, "생산일": "2026-09-03"},
            {"품목": "HOUSING", "생산수량": 80, "불량수량": 2, "생산일": "2026-09-10"},
        ]
        excel = self._analyze(rows)
        self.assertFalse(excel.get("column_mapping", {}).get("equipment", {}).get("column"))
        report, text = self._generate(excel, facts={"total_production": 180, "total_defects": 6, "defect_rate": 3.33})
        self.assertIn("설비별 품질 현황", report["sections_omitted"])
        self.assertIn("설비 정보가 없어", text)
        self.assertNotIn("2호기", text)

    def test_missing_period_omits_comparison(self):
        rows = [{"품목": "BRACKET", "생산수량": 100, "불량수량": 4, "설비": "1호기"}]
        excel = self._analyze(rows)
        report, text = self._generate(excel, facts={"total_production": 100, "total_defects": 4, "defect_rate": 4.0})
        self.assertIn("기간 비교", report["sections_omitted"])
        self.assertIn("기간 비교는 제외", text)
        self.assertNotIn("불량률 변화", text)

    def test_user_path_is_rejected(self):
        excel = self._analyze(_quality_rows())
        result = generate_report({"path": "C:/secret/out.docx"}, _state(excel=excel))
        self.assertFalse(result["ok"])
        self.assertEqual(result["error"], "path_not_allowed")

    def test_report_without_analysis_fails(self):
        result = generate_report({}, _state(execution_results=[]))
        self.assertFalse(result["ok"])
        self.assertEqual(result["error"], "missing_analysis")
        eval_result = evaluate(
            _state(
                execution_results=[
                    {"tool": "generate_report", "ok": False, "detail": result["detail"], "data": result}
                ]
            )
        )
        self.assertEqual(eval_result.status, "fail")
        self.assertTrue(eval_result.user_message or result["detail"])
        self.assertIn("분석 결과가 없어", eval_result.user_message or result["detail"])

    def test_clarification_then_resume_keeps_report_step(self):
        pending = {
            "type": "column",
            "tool": "analyze_excel",
            "document_id": "doc-1",
            "field": "defect_qty",
            "candidates": ["NG_QTY", "검사불량", "최종NG"],
            "include_report": True,
        }
        state = _state("최종NG", execution_results=[], pending_clarification=pending, plan=_report_plan())
        plan = build_work_plan(state)
        self.assertEqual([s.tool for s in plan.steps], ["analyze_excel", "generate_report"])
        self.assertEqual(plan.steps[0].args.get("column_mapping"), {"defect_qty": "최종NG"})
        calls = tool_selector({**state, "plan": plan.model_dump()})
        self.assertEqual([c["tool"] for c in calls["tool_calls"]], ["analyze_excel"])

    def test_excel_clarification_does_not_run_report(self):
        excel = {
            "ok": True,
            "needs_clarification": True,
            "clarification": {"field": "defect_qty", "candidates": ["NG_QTY"], "message": "컬럼을 확인해 주세요."},
        }
        result = evaluate(_state(excel=excel))
        self.assertEqual(result.status, "clarify")
        calls = tool_selector(_state(excel=excel))
        self.assertEqual([c["tool"] for c in calls["tool_calls"]], ["analyze_excel"])
        blocked = generate_report({}, _state(excel=excel))
        self.assertFalse(blocked["ok"])
        self.assertEqual(blocked["error"], "needs_clarification")

    def test_item_defect_rate_does_not_generate_report(self):
        excel = self._analyze(_quality_rows())
        plan = build_work_plan(_state("품목별 불량률 알려줘", execution_results=[], plan={}))
        self.assertEqual([s.tool for s in plan.steps], ["analyze_excel"])
        result = evaluate(
            _state(
                "품목별 불량률 알려줘",
                excel=excel,
                plan={"steps": [{"tool": "analyze_excel", "args": {}}]},
                plan_steps=[{"tool": "analyze_excel"}],
                interpreted_goal="품목별 불량률 알려줘",
            )
        )
        self.assertEqual(result.status, "pass")
        self.assertNotEqual(result.status, "continue")

    def test_executor_dispatches_generate_report(self):
        excel = self._analyze(_quality_rows())
        settings = SimpleNamespace(upload_dir=self.upload_dir)
        with patch("app.tools.report_generator.get_settings", return_value=settings):
            invoked = invoke_tool(
                "generate_report",
                {"depends_on": "analyze_excel"},
                _state(excel=excel, validated_facts={"total_production": 360, "total_defects": 16, "defect_rate": 4.44}),
            )
            executed = tool_executor(
                _state(
                    excel=excel,
                    tool_calls=[{"tool": "generate_report", "args": {"depends_on": "analyze_excel"}}],
                    validated_facts={"total_production": 360, "total_defects": 16, "defect_rate": 4.44},
                )
            )
        self.assertTrue(invoked.ok)
        self.assertEqual(executed["execution_results"][-1]["tool"], "generate_report")
        self.assertTrue(executed["execution_results"][-1]["ok"])

    def test_followup_status_does_not_reuse_completed_report_plan(self):
        plan = build_work_plan(
            _state(
                "2호기 상태 어때?",
                execution_results=[],
                plan=_report_plan(),
                pending_clarification=None,
                interpreted_goal="2호기 상태 분석",
            )
        )
        self.assertEqual([s.tool for s in plan.steps], ["analyze_excel"])
        self.assertNotIn("generate_report", [s.tool for s in plan.steps])
        self.assertEqual(plan.steps[0].args.get("filter"), {"role": "equipment", "value": "2호기"})
        eval_result = evaluate(
            _state(
                "2호기 상태 어때?",
                excel=self._analyze(_quality_rows()),
                plan=plan.model_dump(),
                plan_steps=[s.model_dump() for s in plan.steps],
                interpreted_goal="2호기 상태 분석",
                pending_clarification=None,
            )
        )
        self.assertNotEqual(eval_result.status, "continue")
        self.assertNotIn("generate_report", eval_result.missing_requirements)

    def test_no_cause_hallucination_in_report(self):
        excel = self._analyze(_quality_rows())
        facts = {
            "total_production": 360,
            "total_defects": 16,
            "defect_rate": 4.44,
            "top_item": {"name": "BRACKET", "defect_qty": 14},
            "equipment": {"name": "2호기", "defect_rate": 7.0},
        }
        report, text = self._generate(excel, facts=facts)
        self.assertTrue(report["ok"])
        self.assertIn("2호기", text)
        self.assertIn("BRACKET", text)
        for banned in BANNED_CAUSES:
            self.assertNotIn(banned, text)
        self.assertIn("원인을 확정하기 어렵", text)


if __name__ == "__main__":
    unittest.main()
