from pathlib import Path

from pypdf import PdfReader

TABULAR_SUFFIXES = {".xlsx", ".xls", ".csv"}


def parse_document(path: str | Path, mime: str | None = None) -> dict:
    path = Path(path)
    suffix = path.suffix.lower()
    mime = (mime or "").lower()

    if suffix in TABULAR_SUFFIXES:
        return {
            "text": f"[tabular file: {path.name}]",
            "page_count": 1,
            "file_type": suffix.lstrip("."),
        }
    if suffix in {".hwp", ".hwpx"}:
        return _parse_hwp(path)
    if suffix == ".pdf" or "pdf" in mime:
        return _parse_pdf(path)
    if suffix in {".docx", ".doc"} or "word" in mime or "officedocument" in mime:
        return _parse_docx(path)
    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8", errors="ignore")
        return {"text": text, "page_count": 1, "file_type": "text"}
    raise ValueError(f"지원하지 않는 파일 형식입니다: {path.suffix}")


def _parse_hwp(path: Path) -> dict:
    from app.config import get_settings
    from app.services.hwp_converter import convert_to_pdf, convert_to_text

    settings = get_settings()
    text = (convert_to_text(path, settings.convert_dir) or "").strip()
    if text:
        return {"text": text, "page_count": 1, "file_type": "hwp"}
    pdf_path = convert_to_pdf(path, settings.convert_dir)
    parsed = _parse_pdf(pdf_path)
    parsed["file_type"] = "hwp"
    return parsed
    if suffix in {".docx", ".doc"} or "word" in mime or "officedocument" in mime:
        return _parse_docx(path)
    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8", errors="ignore")
        return {"text": text, "page_count": 1, "file_type": "text"}
    raise ValueError(f"지원하지 않는 파일 형식입니다: {path.suffix}")


def _parse_pdf(path: Path) -> dict:
    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    text = "\n\n".join(pages).strip()
    return {"text": text, "page_count": len(reader.pages), "file_type": "pdf"}


def _parse_docx(path: Path) -> dict:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    return {"text": text, "page_count": 1, "file_type": "docx"}
