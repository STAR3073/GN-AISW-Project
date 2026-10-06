"""검증된 analyze_excel 결과로 제조 품질 보고서를 만든다. 숫자는 재계산하지 않는다."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Mapping
from uuid import uuid4

from app.agent.excel_reply import compact_excel_payload, excel_payload_from_state
from app.config import get_settings
from app.db.database import SessionLocal
from app.models import AgentSession, Attachment, ChatDocument, Document, new_id
from app.services.hwp_converter import kind_from_name

REPORT_TITLE = "주간 품질 현황 보고서"
BANNED_CAUSES = ("설비 노후화", "작업자 실수", "온도 상승", "원자재 문제", "공정조건 이상", "부주의")
REQUIRED_SECTIONS = ("분석 개요", "전체 품질 현황", "주요 확인사항", "추가 확인 필요사항")


def generate_report(args: dict[str, Any] | None, state: Mapping[str, Any] | None) -> dict[str, Any]:
    args = args if isinstance(args, dict) else {}
    state = state or {}
    if args.get("path") or args.get("output_path") or args.get("stored_path"):
        return _fail("서버 임의 경로는 사용할 수 없습니다.", error="path_not_allowed")

    excel = excel_payload_from_state(state)
    if not excel or not excel.get("ok"):
        return _fail("분석 결과가 없어 보고서를 만들 수 없습니다.", error="missing_analysis")
    if excel.get("needs_clarification"):
        return _fail("컬럼 확인이 끝나지 않아 보고서를 만들 수 없습니다.", error="needs_clarification")

    facts = state.get("validated_facts") if isinstance(state.get("validated_facts"), dict) else {}
    model = build_report_model(
        excel=excel,
        facts=facts,
        generated_at=datetime.now(timezone.utc),
    )
    settings = get_settings()
    stored_name = f"{uuid4().hex}.docx"
    dest = settings.upload_dir / stored_name
    try:
        _write_docx(dest, model)
    except Exception as exc:
        return _fail(f"보고서 파일을 만들지 못했습니다: {exc}", error="write_failed")
    if not dest.exists():
        return _fail("보고서 파일이 생성되지 않았습니다.", error="write_failed")

    links = _link_to_session(
        path=dest,
        filename=model["filename"],
        session_id=str(state.get("session_id") or ""),
        preview=model.get("plain_text") or "",
    )
    return {
        "ok": True,
        "error": None,
        "detail": f"보고서 저장: {model['filename']}",
        "title": model["title"],
        "filename": model["filename"],
        "stored_path": str(dest),
        "attachment_id": links.get("attachment_id"),
        "document_id": links.get("document_id"),
        "download_path": links.get("download_path"),
        "sections_included": list(model["sections_included"]),
        "sections_omitted": list(model["sections_omitted"]),
        "used_facts": dict(model["used_facts"]),
        "generated_at": model["generated_at"],
        "file": excel.get("file") or {},
    }


def build_report_model(
    *,
    excel: Mapping[str, Any],
    facts: Mapping[str, Any] | None = None,
    generated_at: datetime | None = None,
) -> dict[str, Any]:
    facts = dict(facts or {})
    compact = compact_excel_payload(excel)
    generated_at = generated_at or datetime.now(timezone.utc)
    stamp = generated_at.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    file_info = excel.get("file") if isinstance(excel.get("file"), dict) else {}
    schema = excel.get("schema") if isinstance(excel.get("schema"), dict) else {}
    quality = excel.get("quality_metrics") if isinstance(excel.get("quality_metrics"), dict) else {}
    groups = excel.get("group_analysis") if isinstance(excel.get("group_analysis"), dict) else {}
    mapping = excel.get("column_mapping") if isinstance(excel.get("column_mapping"), dict) else {}
    change = excel.get("change_detection") if isinstance(excel.get("change_detection"), dict) else {}
    period = excel.get("period_analysis") if isinstance(excel.get("period_analysis"), dict) else {}

    used_facts: dict[str, Any] = {}
    sections: list[dict[str, Any]] = []
    included: list[str] = []
    omitted: list[str] = []

    overview_rows = [
        ("분석 대상 파일", file_info.get("filename") or "-"),
        ("분석 기간", _period_label(change, period)),
        ("분석 데이터 건수", schema.get("row_count") if schema.get("row_count") is not None else "-"),
        ("생성일시", stamp),
    ]
    sections.append({"title": "1. 분석 개요", "kind": "pairs", "rows": overview_rows})
    included.append("분석 개요")

    prod = facts.get("total_production", quality.get("production_qty_total"))
    defect = facts.get("total_defects", quality.get("defect_qty_total"))
    rate = facts.get("defect_rate")
    if rate is None and quality.get("formula") and quality.get("defect_rate_pct") is not None:
        rate = quality.get("defect_rate_pct")
    quality_rows = []
    if prod is not None:
        quality_rows.append(("총 생산수량", prod))
        used_facts["total_production"] = prod
    else:
        quality_rows.append(("총 생산수량", "생산수량 컬럼이 없어 집계하지 않았습니다."))
    if defect is not None:
        quality_rows.append(("총 불량수량", defect))
        used_facts["total_defects"] = defect
    else:
        quality_rows.append(("총 불량수량", "불량수량 컬럼이 없어 집계하지 않았습니다."))
    if prod is not None and defect is not None and rate is not None:
        quality_rows.append(("전체 불량률", f"{rate}%"))
        used_facts["defect_rate"] = rate
        if quality.get("formula"):
            quality_rows.append(("산출 공식", quality.get("formula")))
    else:
        quality_rows.append(("전체 불량률", "생산수량과 불량수량이 없어 계산하지 않았습니다."))
    sections.append({"title": "2. 전체 품질 현황", "kind": "pairs", "rows": quality_rows})
    included.append("전체 품질 현황")

    item_rows = list(groups.get("by_item") or [])
    if item_rows and mapping.get("item", {}).get("column"):
        sections.append(
            {
                "title": "3. 품목별 품질 현황",
                "kind": "table",
                "headers": ["품목", "생산수량", "불량수량", "불량률(%)"],
                "rows": [_group_row(row) for row in item_rows],
                "note": _top_note(item_rows, "품목"),
            }
        )
        included.append("품목별 품질 현황")
    else:
        sections.append(
            {
                "title": "3. 품목별 품질 현황",
                "kind": "note",
                "text": "품목 정보가 없어 품목별 분석은 제외되었습니다.",
            }
        )
        omitted.append("품목별 품질 현황")

    equip_rows = list(groups.get("by_equipment") or [])
    if equip_rows and mapping.get("equipment", {}).get("column"):
        sections.append(
            {
                "title": "4. 설비별 품질 현황",
                "kind": "table",
                "headers": ["설비", "생산수량", "불량수량", "불량률(%)"],
                "rows": [_group_row(row) for row in equip_rows],
                "note": _top_note(equip_rows, "설비"),
            }
        )
        included.append("설비별 품질 현황")
    else:
        sections.append(
            {
                "title": "4. 설비별 품질 현황",
                "kind": "note",
                "text": "설비 정보가 없어 설비별 분석은 제외되었습니다.",
            }
        )
        omitted.append("설비별 품질 현황")

    if _has_period_comparison(facts, change, period):
        period_facts = facts.get("period_comparison") if isinstance(facts.get("period_comparison"), dict) else {}
        prev = period_facts.get("previous_period") or change.get("previous_period")
        latest = period_facts.get("latest_period") or change.get("latest_period")
        rate_delta = period_facts.get("defect_rate_delta_pp", change.get("defect_rate_delta_pp"))
        qty_delta = period_facts.get("defect_qty_delta", change.get("defect_qty_delta"))
        rows = [("현재 기간", latest), ("이전 기간", prev)]
        if rate_delta is not None:
            rows.append(("불량률 변화", f"{rate_delta} %p"))
            used_facts["defect_rate_delta_pp"] = rate_delta
        elif qty_delta is not None:
            rows.append(("불량수량 변화", qty_delta))
        sections.append({"title": "5. 기간 비교", "kind": "pairs", "rows": rows})
        included.append("기간 비교")
    else:
        sections.append(
            {
                "title": "5. 기간 비교",
                "kind": "note",
                "text": "비교할 날짜/기간 정보가 부족하여 기간 비교는 제외되었습니다.",
            }
        )
        omitted.append("기간 비교")

    findings = _findings(facts, item_rows, equip_rows, change)
    sections.append({"title": "6. 주요 확인사항", "kind": "bullets", "items": findings})
    included.append("주요 확인사항")
    sections.append(
        {
            "title": "7. 추가 확인 필요사항",
            "kind": "bullets",
            "items": [
                "현재 데이터만으로 원인을 확정하기 어렵습니다.",
                "설비 상태 및 공정조건 데이터의 추가 확인이 필요합니다.",
            ],
        }
    )
    included.append("추가 확인 필요사항")

    return {
        "title": REPORT_TITLE,
        "filename": f"{REPORT_TITLE}.docx",
        "generated_at": stamp,
        "file": file_info.get("filename") or "",
        "sections": sections,
        "sections_included": included,
        "sections_omitted": omitted,
        "used_facts": used_facts,
        "compact": compact,
        "plain_text": _plain_text(REPORT_TITLE, stamp, file_info.get("filename") or "", sections),
    }


def _fail(detail: str, *, error: str) -> dict[str, Any]:
    return {
        "ok": False,
        "error": error,
        "detail": detail,
        "filename": None,
        "stored_path": None,
        "attachment_id": None,
        "sections_included": [],
        "used_facts": {},
    }


def _period_label(change: Mapping[str, Any], period: Mapping[str, Any]) -> str:
    if change.get("latest_period") and change.get("previous_period"):
        return f"{change.get('previous_period')} ~ {change.get('latest_period')}"
    monthly = list((period or {}).get("monthly") or [])
    if monthly:
        first = monthly[0].get("period")
        last = monthly[-1].get("period")
        if first and last:
            return f"{first} ~ {last}" if first != last else str(last)
    weekly = list((period or {}).get("weekly") or [])
    if weekly:
        first = weekly[0].get("period")
        last = weekly[-1].get("period")
        if first and last:
            return f"{first} ~ {last}" if first != last else str(last)
    return "기간 정보 없음"


def _group_row(row: Mapping[str, Any]) -> list[Any]:
    rate = row.get("defect_rate_pct")
    return [
        row.get("key") or "-",
        row.get("production_qty") if row.get("production_qty") is not None else "-",
        row.get("defect_qty") if row.get("defect_qty") is not None else "-",
        rate if rate is not None else "-",
    ]


def _top_note(rows: list[Mapping[str, Any]], label: str) -> str:
    if not rows:
        return ""
    top = rows[0]
    name = top.get("key")
    if not name:
        return ""
    if top.get("defect_qty") is not None:
        return f"주요 이상 {label}: {name} (불량수량 {top.get('defect_qty')})"
    if top.get("ng_count") is not None:
        return f"주요 이상 {label}: {name} (NG {top.get('ng_count')}건)"
    return f"주요 이상 {label}: {name}"


def _has_period_comparison(facts: Mapping[str, Any], change: Mapping[str, Any], period: Mapping[str, Any]) -> bool:
    period_facts = facts.get("period_comparison") if isinstance(facts.get("period_comparison"), dict) else {}
    if period_facts.get("previous_period") and period_facts.get("latest_period"):
        return any(period_facts.get(k) is not None for k in ("defect_rate_delta_pp", "defect_qty_delta"))
    if change.get("previous_period") and change.get("latest_period"):
        return any(change.get(k) is not None for k in ("defect_rate_delta_pp", "defect_qty_delta", "rows_delta"))
    return False


def _findings(
    facts: Mapping[str, Any],
    items: list[Mapping[str, Any]],
    equipment: list[Mapping[str, Any]],
    change: Mapping[str, Any],
) -> list[str]:
    lines: list[str] = []
    if facts.get("total_production") is not None and facts.get("total_defects") is not None and facts.get("defect_rate") is not None:
        lines.append(
            f"전체 생산수량 {facts['total_production']}, 불량수량 {facts['total_defects']}, 불량률 {facts['defect_rate']}%가 확인되었습니다."
        )
    item = facts.get("top_item") if isinstance(facts.get("top_item"), dict) else (items[0] if items else None)
    if item:
        name = item.get("name") or item.get("key")
        if name:
            extra = f" (불량수량 {item.get('defect_qty')})" if item.get("defect_qty") is not None else ""
            lines.append(f"{name} 품목에서 불량이 가장 많이 확인되었습니다{extra}.")
    equip = facts.get("equipment") if isinstance(facts.get("equipment"), dict) else (equipment[0] if equipment else None)
    if equip:
        name = equip.get("name") or equip.get("key")
        if name:
            rate = equip.get("defect_rate")
            if rate is None:
                rate = equip.get("defect_rate_pct")
            extra = f", 불량률 {rate}%" if rate is not None else ""
            lines.append(f"{name}에서 불량이 확인되었습니다{extra}.")
    period = facts.get("period_comparison") if isinstance(facts.get("period_comparison"), dict) else None
    if period and period.get("previous_period") and period.get("latest_period"):
        delta = period.get("defect_rate_delta_pp")
        if delta is not None:
            direction = "높아졌" if delta > 0 else "낮아졌" if delta < 0 else "같았"
            lines.append(
                f"{period['previous_period']} 대비 {period['latest_period']} 불량률은 {direction}습니다."
            )
    spikes = list((change or {}).get("spikes") or [])
    if spikes:
        spike = spikes[0]
        if spike.get("key"):
            lines.append(f"{spike.get('key')}에서 불량률이 평균보다 높게 나타났습니다.")
    if not lines:
        lines.append("현재 데이터에서 추가로 단정할 수 있는 이상 수치는 제한적입니다.")
    cleaned = []
    for line in lines:
        if any(banned in line for banned in BANNED_CAUSES):
            continue
        cleaned.append(line)
    return cleaned or ["현재 데이터에서 확인된 수치만 보고서에 반영했습니다."]


def _plain_text(title: str, stamp: str, filename: str, sections: list[dict[str, Any]]) -> str:
    lines = [title, f"생성일시: {stamp}", f"분석 대상: {filename or '-'}", ""]
    for section in sections:
        lines.append(section.get("title") or "")
        if section.get("kind") == "pairs":
            for key, value in section.get("rows") or []:
                lines.append(f"- {key}: {value}")
        elif section.get("kind") == "table":
            for row in section.get("rows") or []:
                lines.append("- " + " / ".join(str(v) for v in row))
            if section.get("note"):
                lines.append(str(section["note"]))
        elif section.get("kind") == "note":
            lines.append(str(section.get("text") or ""))
        elif section.get("kind") == "bullets":
            for item in section.get("items") or []:
                lines.append(f"- {item}")
        lines.append("")
    return "\n".join(lines).strip()


def _write_docx(path: Path, model: Mapping[str, Any]) -> None:
    from docx import Document as DocxDocument
    from docx.shared import Pt

    doc = DocxDocument()
    doc.add_heading(str(model.get("title") or REPORT_TITLE), level=0)
    meta = doc.add_paragraph()
    meta.add_run(f"생성일시: {model.get('generated_at') or ''}").font.size = Pt(11)
    doc.add_paragraph(f"분석 대상: {model.get('file') or '-'}")

    for section in model.get("sections") or []:
        doc.add_heading(str(section.get("title") or ""), level=1)
        kind = section.get("kind")
        if kind == "pairs":
            table = doc.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            hdr[0].text = "항목"
            hdr[1].text = "내용"
            for key, value in section.get("rows") or []:
                cells = table.add_row().cells
                cells[0].text = str(key)
                cells[1].text = "-" if value is None else str(value)
        elif kind == "table":
            headers = list(section.get("headers") or [])
            rows = list(section.get("rows") or [])
            table = doc.add_table(rows=1, cols=max(len(headers), 1))
            table.style = "Table Grid"
            for idx, header in enumerate(headers):
                table.rows[0].cells[idx].text = str(header)
            for row in rows:
                cells = table.add_row().cells
                for idx, value in enumerate(row):
                    if idx < len(cells):
                        cells[idx].text = "-" if value is None else str(value)
            if section.get("note"):
                doc.add_paragraph(str(section["note"]))
        elif kind == "note":
            doc.add_paragraph(str(section.get("text") or ""))
        elif kind == "bullets":
            for item in section.get("items") or []:
                doc.add_paragraph(str(item), style="List Bullet")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def _link_to_session(*, path: Path, filename: str, session_id: str, preview: str) -> dict[str, Any]:
    out = {"attachment_id": None, "document_id": None, "download_path": None}
    if not session_id:
        return out
    db = SessionLocal()
    try:
        session = db.get(AgentSession, session_id)
        if not session:
            return out
        size = path.stat().st_size if path.exists() else 0
        att = Attachment(
            id=new_id(),
            filename=filename,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            kind=kind_from_name(filename),
            stored_path=str(path),
            size_bytes=size,
        )
        db.add(att)
        db.flush()
        doc = Document(
            id=new_id(),
            project_id=session.project_id,
            filename=filename,
            mime=att.mime,
            stored_path=str(path),
            raw_text=(preview or "")[:8000],
            page_count=1,
            size_bytes=size,
        )
        db.add(doc)
        db.flush()
        link = (
            db.query(ChatDocument)
            .filter(ChatDocument.session_id == session.id, ChatDocument.document_id == doc.id)
            .first()
        )
        if not link:
            db.add(ChatDocument(session_id=session.id, document_id=doc.id))
        att.source_id = doc.id
        db.commit()
        out["attachment_id"] = att.id
        out["document_id"] = doc.id
        out["download_path"] = f"/api/attachments/{att.id}/file?download=1"
        return out
    except Exception:
        db.rollback()
        return out
    finally:
        db.close()
